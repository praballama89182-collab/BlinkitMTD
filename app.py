import os
import io
import pandas as pd
import numpy as np
import streamlit as st
import plotly.graph_objects as go
from plotly.subplots import make_subplots
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

st.set_page_config(
    page_title="Blinkit Ad Report & Sales Merger & Analytics",
    page_icon="data:image/svg+xml;base64,PHN2ZyB4bWxucz0iaHR0cDovL3d3dy53My5vcmcvMjAwMC9zdmciIHZpZXdCb3g9IjAgMCA1MTIgNTEyIj48cmVjdCB3aWR0aD0iNTEyIiBoZWlnaHQ9IjUxMiIgcng9IjExMiIgZmlsbD0iI2ZkZDA0MyIvPjx0ZXh0IHg9IjI1NiIgeT0iMzIwIiBmb250LWZhbWlseT0iQXJpYWwsIHNhbnMtc2VyaWYiIGZvbnQtd2VpZ2h0PSJib2x0IiBmb250LXNpemU9IjIxMCIgdGV4dC1hbmNob3I9Im1pZGRsZSI+PGZzcGFuIGZpbGw9IiMwMDAwMDAiPmJsaW5rPC9mc3Bhbj48ZnNwYW4gZmlsbD0iIzAwODEwNyI+aXQ8L2ZzcGFuPjwvdGV4dD48L3N2Zz4=",
    layout="wide"
)

st.markdown("""
<style>
    :root {
        --dash-blue: #4A90C2;
        --dash-blue-dark: #285B7A;
        --dash-blue-light: #EEF6FC;
        --dash-blue-soft: #DCECF8;
        --dash-border: #C8DCEB;
        --dash-text: #29465B;
    }

    .stApp {
        background: linear-gradient(180deg, #F6FAFD 0%, #FFFFFF 38%);
        color: var(--dash-text);
    }

    [data-testid="stHeader"] {
        background: rgba(255,255,255,0.92);
    }

    h1, h2, h3 {
        color: var(--dash-blue-dark) !important;
        font-weight: 700 !important;
    }

    [data-testid="stMetric"] {
        background: linear-gradient(135deg, #FFFFFF 0%, #EEF6FF 100%);
        border: 1px solid var(--dash-border);
        border-left: 5px solid #6AA6CF;
        border-radius: 12px;
        padding: 14px 18px;
        box-shadow: 0 3px 12px rgba(74, 144, 194, 0.10);
    }

    [data-testid="stMetricLabel"] {
        color: #52708D !important;
        font-weight: 600 !important;
    }

    [data-testid="stMetricValue"] {
        color: var(--dash-blue-dark) !important;
        font-weight: 750 !important;
    }

    .stTabs [data-baseweb="tab-list"] {
        gap: 5px;
        background: #F0F7FC;
        border: 1px solid var(--dash-border);
        border-radius: 12px;
        padding: 5px;
    }

    .stTabs [data-baseweb="tab"] {
        height: 42px;
        border-radius: 7px;
        color: #315A7D;
        font-weight: 600;
    }

    .stTabs [aria-selected="true"] {
        background: linear-gradient(135deg, #6AA6CF, #4A90C2) !important;
        color: #FFFFFF !important;
        box-shadow: 0 2px 7px rgba(74, 144, 194, 0.20);
    }

    [data-testid="stFileUploader"] {
        background: #FFFFFF;
        border: 1px solid var(--dash-border);
        border-radius: 12px;
        padding: 8px;
        box-shadow: 0 2px 8px rgba(74, 144, 194, 0.07);
    }

    [data-testid="stFileUploaderDropzone"] {
        background: #F7FBFE;
        border: 1px dashed #9FC4DE;
        border-radius: 9px;
    }

    [data-testid="stFileUploader"] button,
    .stDownloadButton button,
    .stButton button {
        background: linear-gradient(135deg, #6AA6CF, #4A90C2) !important;
        color: #FFFFFF !important;
        border: 0 !important;
        border-radius: 8px !important;
        font-weight: 650 !important;
        box-shadow: 0 2px 7px rgba(11, 92, 173, 0.20);
    }

    [data-testid="stFileUploader"] button:hover,
    .stDownloadButton button:hover,
    .stButton button:hover {
        background: #397BAA !important;
        color: #FFFFFF !important;
    }

    [data-testid="stDataFrame"] {
        border: 1px solid var(--dash-border);
        border-radius: 10px;
        overflow: hidden;
        box-shadow: 0 2px 9px rgba(11, 92, 173, 0.06);
    }

    [data-testid="stDataFrame"] th,
    [data-testid="stDataFrame"] td {
        text-align: center !important;
        vertical-align: middle !important;
    }

    [data-testid="stDataFrame"] th {
        background: #E5F1F9 !important;
        color: #285B7A !important;
        font-weight: 700 !important;
    }

    [data-baseweb="select"] > div {
        border-color: #AFC7DF !important;
        border-radius: 8px !important;
    }

    [data-testid="stRadio"] label,
    [data-testid="stSelectbox"] label,
    [data-testid="stMultiSelect"] label {
        color: #315A7D !important;
        font-weight: 600 !important;
    }

    hr {
        border-color: #D7E4F2 !important;
    }
</style>
""", unsafe_allow_html=True)

st.markdown("""
<div style="display: flex; align-items: center; gap: 14px; margin-bottom: 5px;">
    <svg width="48" height="48" viewBox="0 0 512 512" style="border-radius: 12px; box-shadow: 0 2px 6px rgba(0,0,0,0.15);">
        <rect width="512" height="512" rx="112" fill="#fdd043"/>
        <text x="256" y="320" font-family="Arial, sans-serif" font-weight="bold" font-size="210" text-anchor="middle">
            <tspan fill="#000000">blink</tspan><tspan fill="#008107">it</tspan>
        </text>
    </svg>
    <h1 style="margin: 0; padding: 0; line-height: 1.1;">Blinkit Ad Report & Sales Merger & Analytics</h1>
</div>
""", unsafe_allow_html=True)

st.write("Upload up to 5 monthly Ad Campaign spreadsheets and up to 5 Sales Reports (.xlsx, .xls) to analyze TACOS, TROAS, Overall Sales, Orders, and multi-metric trends.")

def style_and_export_pivot(pivot_df, sheet_name="Comparison"):
    buffer = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    
    top_header_fill = PatternFill(start_color="4A90C2", end_color="4A90C2", fill_type="solid")
    top_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    sec_header_fill = PatternFill(start_color="E5F1F9", end_color="E5F1F9", fill_type="solid")
    sec_header_font = Font(name="Calibri", size=10, bold=True, color="285B7A")
    index_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    index_font = Font(name="Calibri", size=10, bold=True, color="000000")
    total_fill = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")
    total_font = Font(name="Calibri", size=10, bold=True, color="000000")
    data_font = Font(name="Calibri", size=10, color="000000")
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=False)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=False)
    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF')
    )

    if isinstance(pivot_df.columns, pd.MultiIndex):
        metrics = pivot_df.columns.get_level_values(1)
        entity_title = pivot_df.index.name if pivot_df.index.name else ""

        ws.cell(row=1, column=1, value="").fill = top_header_fill
        ws.cell(row=1, column=1).border = thin_border
        
        current_col = 2
        for month in pivot_df.columns.get_level_values(0).unique():
            sub_cols = [c for c in pivot_df.columns if c[0] == month]
            num_sub = len(sub_cols)
            if num_sub > 1:
                ws.merge_cells(start_row=1, start_column=current_col, end_row=1, end_column=current_col + num_sub - 1)
            cell = ws.cell(row=1, column=current_col, value=str(month))
            cell.fill = top_header_fill
            cell.font = top_header_font
            cell.alignment = align_center
            
            for c in range(current_col, current_col + num_sub):
                ws.cell(row=1, column=c).border = thin_border
                ws.cell(row=1, column=c).fill = top_header_fill
            current_col += num_sub

        cell_a2 = ws.cell(row=2, column=1, value=str(entity_title))
        cell_a2.fill = index_fill
        cell_a2.font = index_font
        cell_a2.alignment = align_center
        cell_a2.border = thin_border

        for col_idx, metric in enumerate(metrics, start=2):
            cell = ws.cell(row=2, column=col_idx, value=str(metric))
            cell.fill = sec_header_fill
            cell.font = sec_header_font
            cell.alignment = align_center
            cell.border = thin_border

        start_data_row = 3
        for r_idx, (idx_val, row_data) in enumerate(pivot_df.iterrows(), start=start_data_row):
            is_grand_total = (str(idx_val).strip().lower() == 'grand total')
            
            idx_cell = ws.cell(row=r_idx, column=1, value=str(idx_val))
            idx_cell.fill = total_fill if is_grand_total else index_fill
            idx_cell.font = total_font if is_grand_total else index_font
            idx_cell.alignment = align_left
            idx_cell.border = thin_border

            for c_idx, (m_col, val) in enumerate(zip(metrics, row_data), start=2):
                val_cell = ws.cell(row=r_idx, column=c_idx)
                val_cell.font = total_font if is_grand_total else data_font
                if is_grand_total:
                    val_cell.fill = total_fill
                val_cell.alignment = align_center
                val_cell.border = thin_border

                if isinstance(val, (int, float, np.number)):
                    if str(m_col).upper() in ['ACOS', 'TACOS']:
                        val_cell.value = val / 100.0 if val > 1 else val
                        val_cell.number_format = '0.00%'
                    else:
                        val_cell.value = round(val, 2)
                        val_cell.number_format = '#,##0.00' if not float(val).is_integer() else '#,##0'
                else:
                    val_cell.value = val
    else:
        for col_idx, col_name in enumerate(pivot_df.columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=str(col_name))
            cell.fill = top_header_fill
            cell.font = top_header_font
            cell.alignment = align_center
            cell.border = thin_border

        for r_idx, (idx_val, row_data) in enumerate(pivot_df.iterrows(), start=2):
            is_pct_row = (str(row_data.iloc[0]).strip().lower() == 'percentage %')
            
            for c_idx, (col_name, val) in enumerate(zip(pivot_df.columns, row_data), start=1):
                val_cell = ws.cell(row=r_idx, column=c_idx)
                val_cell.font = total_font if is_pct_row else data_font
                val_cell.alignment = align_center
                val_cell.border = thin_border
                
                if str(col_name).upper() in ['ACOS', 'TACOS'] and not is_pct_row and isinstance(val, (int, float, np.number)):
                    val_cell.value = val / 100.0 if val > 1 else val
                    val_cell.number_format = '0.00%'
                elif isinstance(val, float):
                    val_cell.value = round(val, 2)
                    val_cell.number_format = '#,##0.00'
                elif isinstance(val, (int, np.integer)):
                    val_cell.value = val
                    val_cell.number_format = '#,##0'
                else:
                    val_cell.value = val

                if is_pct_row and c_idx > 1:
                    val_str = str(val).replace('%', '').replace('+', '').strip()
                    try:
                        num_v = float(val_str)
                        if num_v < 0:
                            val_cell.fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
                            val_cell.font = Font(name="Calibri", size=10, bold=True, color="721C24")
                        elif num_v > 0:
                            val_cell.fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
                            val_cell.font = Font(name="Calibri", size=10, bold=True, color="155724")
                    except ValueError:
                        pass
                elif is_pct_row:
                    val_cell.fill = sec_header_fill

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

st.markdown("### 📤 Upload Ad Campaign Reports (Up to 5)")
col1, col2, col3, col4, col5 = st.columns(5)
uploaded_files = []

with col1:
    f1 = st.file_uploader("Upload Ad File 1", type=["xlsx", "xls", "xlsb", "xlsm"], key="file1")
    if f1: uploaded_files.append(f1)
with col2:
    f2 = st.file_uploader("Upload Ad File 2", type=["xlsx", "xls", "xlsb", "xlsm"], key="file2")
    if f2: uploaded_files.append(f2)
with col3:
    f3 = st.file_uploader("Upload Ad File 3", type=["xlsx", "xls", "xlsb", "xlsm"], key="file3")
    if f3: uploaded_files.append(f3)
with col4:
    f4 = st.file_uploader("Upload Ad File 4", type=["xlsx", "xls", "xlsb", "xlsm"], key="file4")
    if f4: uploaded_files.append(f4)
with col5:
    f5 = st.file_uploader("Upload Ad File 5", type=["xlsx", "xls", "xlsb", "xlsm"], key="file5")
    if f5: uploaded_files.append(f5)

st.markdown("### 📥 Upload Sales Reports (Up to 5)")
scol1, scol2, scol3, scol4, scol5 = st.columns(5)
uploaded_sales_files = []

with scol1:
    s1 = st.file_uploader("Upload Sales File 1", type=["xlsx", "xls"], key="sales1")
    if s1: uploaded_sales_files.append(s1)
with scol2:
    s2 = st.file_uploader("Upload Sales File 2", type=["xlsx", "xls"], key="sales2")
    if s2: uploaded_sales_files.append(s2)
with scol3:
    s3 = st.file_uploader("Upload Sales File 3", type=["xlsx", "xls"], key="sales3")
    if s3: uploaded_sales_files.append(s3)
with scol4:
    s4 = st.file_uploader("Upload Sales File 4", type=["xlsx", "xls"], key="sales4")
    if s4: uploaded_sales_files.append(s4)
with scol5:
    s5 = st.file_uploader("Upload Sales File 5", type=["xlsx", "xls"], key="sales5")
    if s5: uploaded_sales_files.append(s5)

# Process Sales Reports
sales_dfs_list = []
if uploaded_sales_files:
    for s_file in uploaded_sales_files:
        try:
            s_xls = pd.ExcelFile(s_file)
            for s_sheet in s_xls.sheet_names:
                sdf = pd.read_excel(s_xls, sheet_name=s_sheet)
                sales_dfs_list.append(sdf)
        except Exception as e:
            st.error(f"Error reading sales file {s_file.name}: {str(e)}")

master_sales_df = pd.DataFrame()
if sales_dfs_list:
    master_sales_df = pd.concat(sales_dfs_list, ignore_index=True)
    date_cand = next((c for c in ['Order Date', 'date', 'Date', 'ORDER_DATE'] if c in master_sales_df.columns), None)
    gross_cand = next((c for c in ['Total Gross Bill Amount', 'Gross Amount', 'Total Amount', 'Bill Amount'] if c in master_sales_df.columns), None)
            
    if date_cand:
        master_sales_df['_sales_dt'] = pd.to_datetime(master_sales_df[date_cand], errors='coerce')
        master_sales_df['Month'] = master_sales_df['_sales_dt'].dt.strftime('%B').str.upper()
        master_sales_df['Date_str'] = master_sales_df['_sales_dt'].dt.strftime('%Y-%m-%d')
        master_sales_df['Day'] = master_sales_df['_sales_dt'].dt.day
    if gross_cand:
        master_sales_df['_gross_amt'] = pd.to_numeric(master_sales_df[gross_cand], errors='coerce').fillna(0.0)
    else:
        master_sales_df['_gross_amt'] = 0.0

if uploaded_files:
    consolidated_dfs = []
    raw_files_dict = {}
    consolidated_raw_tabs = {}

    for uploaded_file in uploaded_files:
        fallback_month_name = os.path.splitext(uploaded_file.name)[0].upper()
        raw_files_dict[uploaded_file.name] = {}
        
        try:
            xls = pd.ExcelFile(uploaded_file)
            sheets_to_process = xls.sheet_names[1:] if len(xls.sheet_names) > 1 else xls.sheet_names
            valid_sheet_names = [s for s in sheets_to_process if s.strip().lower() not in ['consolidated_master', 'summary', 'overall']]

            for sheet_name in valid_sheet_names:
                df = pd.read_excel(xls, sheet_name=sheet_name)
                raw_files_dict[uploaded_file.name][sheet_name] = df.copy()
                
                date_col = next((c for c in ['Date', 'date', 'Day', 'DATE', 'date_ist', 'timestamp'] if c in df.columns), None)

                if date_col:
                    dt_series = pd.to_datetime(df[date_col], format='mixed', dayfirst=False, errors='coerce')
                    month_series = dt_series.dt.strftime('%B').str.upper()
                    df['Month'] = month_series.fillna(fallback_month_name)
                    df['_date_dt'] = dt_series
                    df['Date'] = dt_series.dt.strftime('%Y-%m-%d')
                else:
                    df['Month'] = fallback_month_name
                    df['_date_dt'] = pd.NaT
                    df['Date'] = np.nan

                df_raw = df.copy()
                if sheet_name not in consolidated_raw_tabs:
                    consolidated_raw_tabs[sheet_name] = []
                consolidated_raw_tabs[sheet_name].append(df_raw)
                
                df_consolidated = df.copy()
                if 'Month' in df_consolidated.columns:
                    col_month = df_consolidated.pop('Month')
                    df_consolidated.insert(0, 'Month', col_month)
                else:
                    df_consolidated.insert(0, 'Month', fallback_month_name)
                    
                df_consolidated.insert(1, 'Tab Name', sheet_name)
                
                if 'Date' in df_consolidated.columns:
                    col_date = df_consolidated.pop('Date')
                    df_consolidated.insert(2, 'Date', col_date)
                else:
                    df_consolidated.insert(2, 'Date', np.nan)

                consolidated_dfs.append(df_consolidated)
        except Exception as e:
            st.error(f"Error processing {uploaded_file.name}: {str(e)}")

    if consolidated_dfs:
        final_df = pd.concat(consolidated_dfs, ignore_index=True)
        
        base_cols = ['Month', 'Tab Name', 'Date']
        product_listing_cols = []
        for df in consolidated_dfs:
            if 'PRODUCT_LISTING' in str(df['Tab Name'].values).upper():
                pl_df = df[df['Tab Name'].str.upper() == 'PRODUCT_LISTING']
                product_listing_cols = [c for c in pl_df.columns if c not in base_cols]
                break
        
        remaining_cols = [c for c in final_df.columns if c not in base_cols and c not in product_listing_cols]
        final_df = final_df.reindex(columns=base_cols + product_listing_cols + remaining_cols)

        match_col = 'Match Type' if 'Match Type' in final_df.columns else None
        target_col = 'Targeting Type' if 'Targeting Type' in final_df.columns else None
        tab_col = 'Tab Name' if 'Tab Name' in final_df.columns else None

        def is_na_series(series):
            if series is None or series.empty:
                return pd.Series(True, index=final_df.index)
            cleaned = series.astype(str).str.strip().str.upper()
            return series.isna() | cleaned.isin(['NA', 'N/A', 'NAN', 'NONE', ''])

        if match_col:
            na_match = is_na_series(final_df[match_col])
            if target_col:
                valid_target = ~is_na_series(final_df[target_col])
                fill_from_target = na_match & valid_target
                final_df.loc[fill_from_target, match_col] = final_df.loc[fill_from_target, target_col]
                na_match = is_na_series(final_df[match_col])

            if tab_col:
                final_df.loc[na_match, match_col] = final_df.loc[na_match, tab_col]
        elif target_col:
            final_df['Match Type'] = final_df[target_col]
            na_match = is_na_series(final_df['Match Type'])
            if tab_col:
                final_df.loc[na_match, 'Match Type'] = final_df.loc[na_match, tab_col]
            match_col = 'Match Type'

        def parse_robust_numeric(series):
            if series is None:
                return pd.Series(0.0, index=[])
            s_str = series.astype(str).str.replace(',', '', regex=False)
            return pd.to_numeric(s_str, errors='coerce').fillna(0.0)

        def get_numeric_col(df, possible_cols):
            for col in possible_cols:
                if col in df.columns:
                    return parse_robust_numeric(df[col])
            return pd.Series(0.0, index=df.index)

        final_df['_impressions'] = get_numeric_col(final_df, ['Impressions'])
        final_df['_direct_atc'] = get_numeric_col(final_df, ['Direct ATC'])
        final_df['_indirect_atc'] = get_numeric_col(final_df, ['Indirect ATC'])
        final_df['_atc'] = final_df['_direct_atc'] + final_df['_indirect_atc']

        final_df['_direct_orders'] = get_numeric_col(final_df, ['Direct Quantities Sold', 'Direct Orders'])
        final_df['_indirect_orders'] = get_numeric_col(final_df, ['Indirect Quantities Sold', 'Indirect Orders'])
        final_df['_orders'] = final_df['_direct_orders'] + final_df['_indirect_orders']

        final_df['_direct_sales'] = get_numeric_col(final_df, ['Direct Sales'])
        final_df['_indirect_sales'] = get_numeric_col(final_df, ['Indirect Sales'])
        final_df['_sales'] = final_df['_direct_sales'] + final_df['_indirect_sales']

        final_df['_budget_consumed'] = get_numeric_col(final_df, ['Estimated Budget Consumed', 'Budget Consumed', 'Spend'])

        if match_col and match_col in final_df.columns:
            final_df['Ad Type Combined'] = final_df[match_col].fillna("Other").astype(str).str.title()
        else:
            final_df['Ad Type Combined'] = "Other"

        date_col = next((c for c in ['Date', 'date', 'Day', 'DATE', 'date_ist', 'timestamp'] if c in final_df.columns), None)

        if date_col and '_date_dt' in final_df.columns:
            def assign_week_formatted(row):
                dt = row['_date_dt']
                if pd.isna(dt): return np.nan
                day = dt.day
                if 1 <= day <= 7: return "Week 1 (1-7)"
                elif 8 <= day <= 14: return "Week 2 (8-14)"
                elif 15 <= day <= 21: return "Week 3 (15-21)"
                elif 22 <= day <= 28: return "Week 4 (22-28)"
                elif day >= 29: return "Week 5 (29-31)"
                return np.nan
            final_df['Week'] = final_df.apply(assign_week_formatted, axis=1)
        else:
            final_df['Week'] = np.nan

        def compute_grouped_table(df_subset, group_col, selected_item="All", search_query=""):
            if group_col not in df_subset.columns:
                return pd.DataFrame()
            
            df_working = df_subset.dropna(subset=[group_col]).copy()
            df_working[group_col] = df_working[group_col].astype(str)
            
            if selected_item and selected_item != "All":
                df_working = df_working[df_working[group_col] == selected_item]

            if search_query and search_query.strip():
                df_working = df_working[df_working[group_col].str.contains(search_query.strip(), case=False, na=False)]
            
            if df_working.empty:
                return pd.DataFrame()

            grouped = df_working.groupby(group_col).agg(
                IMPRESSIONS=('_impressions', 'sum'),
                ATC=('_atc', 'sum'),
                ORDERS=('_orders', 'sum'),
                SPENDS=('_budget_consumed', 'sum'),
                SALES=('_sales', 'sum')
            ).reset_index()

            grouped['CPM'] = grouped.apply(lambda r: round((r['SPENDS'] / r['IMPRESSIONS']) * 1000, 2) if r['IMPRESSIONS'] > 0 else 0.0, axis=1)
            grouped['ROAS'] = grouped.apply(lambda r: round(r['SALES'] / r['SPENDS'], 2) if r['SPENDS'] > 0 else 0.0, axis=1)
            grouped['ACOS'] = grouped.apply(lambda r: round((r['SPENDS'] / r['SALES']) * 100, 2) if r['SALES'] > 0 else 0.0, axis=1)

            display_name = group_col.upper()
            if group_col == 'Campaign Name': display_name = 'CAMPAIGN NAME'
            elif group_col == 'Ad Type Combined': display_name = 'MATCH / AD TYPE'

            grouped = grouped.rename(columns={group_col: display_name})
            grouped = grouped.sort_values(by='SPENDS', ascending=False).reset_index(drop=True)

            col_order = [display_name, 'IMPRESSIONS', 'CPM', 'ATC', 'ORDERS', 'SPENDS', 'SALES', 'ROAS', 'ACOS']
            res_df = grouped.reindex(columns=col_order)
            res_df['SPENDS'] = res_df['SPENDS'].round(2)
            res_df['SALES'] = res_df['SALES'].round(2)
            return res_df

        MONTH_ORDER_ASC = {
            "JANUARY": 1, "FEBRUARY": 2, "MARCH": 3, "APRIL": 4,
            "MAY": 5, "JUNE": 6, "JULY": 7, "AUGUST": 8,
            "SEPTEMBER": 9, "OCTOBER": 10, "NOVEMBER": 11, "DECEMBER": 12
        }

        def get_calendar_months_asc(df_input):
            months = [str(x).strip().upper() for x in df_input['Month'].dropna().unique()]
            return sorted(months, key=lambda x: MONTH_ORDER_ASC.get(x, 99))

        def create_monthly_summary_table(df_input, selected_months=None):
            working_df = df_input.dropna(subset=['Month']).copy()
            if selected_months:
                working_df = working_df[working_df['Month'].isin(selected_months)]

            if working_df.empty:
                return pd.DataFrame()

            all_months = get_calendar_months_asc(working_df)
            has_sales = not master_sales_df.empty and '_sales_dt' in master_sales_df.columns

            rows_list = []
            for m in all_months:
                m_df = working_df[working_df['Month'] == m]
                imp = m_df['_impressions'].sum()
                atc = m_df['_atc'].sum()
                ords = m_df['_orders'].sum()
                spends = round(m_df['_budget_consumed'].sum(), 2)
                sales = round(m_df['_sales'].sum(), 2)
                cpm = round((spends / imp) * 1000, 2) if imp > 0 else 0.0
                roas = round(sales / spends, 2) if spends > 0 else 0.0
                acos = round((spends / sales) * 100, 2) if sales > 0 else 0.0
                
                row_dict = {
                    'Month': m,
                    'Impressions': imp,
                    'CPM': cpm,
                    'ATC': atc,
                    'Orders': ords,
                }
                
                if has_sales:
                    m_sales_df = master_sales_df[master_sales_df['Month'] == m]
                    ov_sales = round(m_sales_df['_gross_amt'].sum(), 2)
                    ov_orders = len(m_sales_df)
                    tacos = round((spends / ov_sales) * 100, 2) if ov_sales > 0 else 0.0
                    troas = round(ov_sales / spends, 2) if spends > 0 else 0.0
                    
                    row_dict['Overall Sales'] = ov_sales
                    row_dict['Overall Orders'] = ov_orders
                    row_dict['TACOS'] = f"{tacos:.2f}%"
                    row_dict['TROAS'] = troas

                row_dict.update({
                    'Spends': spends,
                    'Sales': sales,
                    'ROAS': roas,
                    'ACOS': f"{acos:.2f}%"
                })
                rows_list.append(row_dict)

            return pd.DataFrame(rows_list)

        def create_mom_comparison_table(df_input, entity_col, selected_months=None):
            if entity_col not in df_input.columns:
                return pd.DataFrame()
            
            working_df = df_input.dropna(subset=[entity_col, 'Month']).copy()
            if selected_months:
                working_df = working_df[working_df['Month'].isin(selected_months)]

            if working_df.empty:
                return pd.DataFrame()
            
            grouped = working_df.groupby([entity_col, 'Month']).agg(
                Impressions=('_impressions', 'sum'),
                ATC=('_atc', 'sum'),
                Orders=('_orders', 'sum'),
                Spends=('_budget_consumed', 'sum'),
                Sales=('_sales', 'sum')
            ).reset_index()

            grouped['CPM'] = grouped.apply(lambda r: round((r['Spends'] / r['Impressions']) * 1000, 2) if r['Impressions'] > 0 else 0.0, axis=1)
            grouped['ROAS'] = grouped.apply(lambda r: round(r['Sales'] / r['Spends'], 2) if r['Spends'] > 0 else 0.0, axis=1)
            grouped['ACOS'] = grouped.apply(lambda r: round((r['Spends'] / r['Sales']) * 100, 2) if r['Sales'] > 0 else 0.0, axis=1)

            pivot_df = grouped.pivot(
                index=entity_col,
                columns='Month',
                values=['Impressions', 'CPM', 'ATC', 'Orders', 'Spends', 'Sales', 'ROAS', 'ACOS']
            )

            metrics_order = ['Impressions', 'CPM', 'ATC', 'Orders', 'Spends', 'Sales', 'ROAS', 'ACOS']
            all_months_asc = get_calendar_months_asc(working_df)
            all_months_sorted = sorted(all_months_asc, key=lambda x: MONTH_ORDER_ASC.get(str(x).strip().upper(), 99), reverse=False)
            
            pivot_df = pivot_df.reorder_levels([1, 0], axis=1)
            sorted_cols = pd.MultiIndex.from_product([all_months_sorted, metrics_order], names=['Month', 'Metric'])
            pivot_df = pivot_df.reindex(columns=sorted_cols).fillna(0)

            total_spends_per_entity = grouped.groupby(entity_col)['Spends'].sum()
            sorted_entities = total_spends_per_entity.sort_values(ascending=False).index
            pivot_df = pivot_df.reindex(sorted_entities)

            grand_total_series = {}
            for month in all_months_sorted:
                month_df = working_df[working_df['Month'] == month]
                total_imp = month_df['_impressions'].sum()
                total_atc = month_df['_atc'].sum()
                total_orders = month_df['_orders'].sum()
                total_spends = month_df['_budget_consumed'].sum()
                total_sales = month_df['_sales'].sum()

                grand_total_series[(month, 'Impressions')] = total_imp
                grand_total_series[(month, 'CPM')] = round((total_spends / total_imp) * 1000, 2) if total_imp > 0 else 0.0
                grand_total_series[(month, 'ATC')] = total_atc
                grand_total_series[(month, 'Orders')] = total_orders
                grand_total_series[(month, 'Spends')] = round(total_spends, 2)
                grand_total_series[(month, 'Sales')] = round(total_sales, 2)
                grand_total_series[(month, 'ROAS')] = round(total_sales / total_spends, 2) if total_spends > 0 else 0.0
                grand_total_series[(month, 'ACOS')] = round((total_spends / total_sales) * 100, 2) if total_sales > 0 else 0.0

            pivot_df.loc['Grand Total'] = grand_total_series
            pivot_df.index.name = entity_col.upper()
            return pivot_df

        def create_weekly_comparison_table_with_sales(df_input, selected_months=None):
            working_df = df_input.dropna(subset=['Week', 'Month']).copy()
            if selected_months:
                working_df = working_df[working_df['Month'].isin(selected_months)]
            if working_df.empty:
                return pd.DataFrame()

            weeks_list = sorted(working_df['Week'].dropna().unique(), key=lambda w: int(''.join(filter(str.isdigit, str(w))) or 99))
            has_sales = not master_sales_df.empty and '_sales_dt' in master_sales_df.columns

            rows = []
            for w in weeks_list:
                w_df = working_df[working_df['Week'] == w]
                spends = round(w_df['_budget_consumed'].sum(), 2)
                sales = round(w_df['_sales'].sum(), 2)
                orders = int(w_df['_orders'].sum())
                ad_roas = round(sales / spends, 2) if spends > 0 else 0.0
                
                row_data = {
                    'Week': w,
                    'Spends': spends,
                    'Ad Sales': sales,
                    'Ad ROAS': ad_roas,
                    'Ad Orders': orders
                }
                
                if has_sales:
                    w_dates = w_df['Date'].dropna().unique()
                    w_sales_df = master_sales_df[master_sales_df['Date_str'].isin(w_dates)]
                    ov_sales = round(w_sales_df['_gross_amt'].sum(), 2)
                    ov_orders = len(w_sales_df)
                    tacos = round((spends / ov_sales) * 100, 2) if ov_sales > 0 else 0.0
                    troas = round(ov_sales / spends, 2) if spends > 0 else 0.0
                    
                    row_data['Overall Sales'] = ov_sales
                    row_data['Overall Orders'] = ov_orders
                    row_data['TACOS'] = f"{tacos:.2f}%"
                    row_data['TROAS'] = troas

                rows.append(row_data)
            return pd.DataFrame(rows)

        def create_daily_tracker_table(selected_month):
            if master_sales_df.empty or '_sales_dt' not in master_sales_df.columns:
                return pd.DataFrame()

            m_df = master_sales_df[master_sales_df['Month'] == selected_month].copy()
            if m_df.empty:
                return pd.DataFrame()

            ad_df = final_df[final_df['Month'] == selected_month].copy() if 'Month' in final_df.columns else pd.DataFrame()

            rows = []
            for day_i in range(1, 32):
                day_sales = m_df[m_df['Day'] == day_i]
                ov_sales = round(day_sales['_gross_amt'].sum(), 2)
                ov_orders = len(day_sales)
                
                day_spends = 0.0
                if not ad_df.empty and 'Date' in ad_df.columns:
                    matched_ad = ad_df[ad_df['_date_dt'].dt.day == day_i]
                    day_spends = round(matched_ad['_budget_consumed'].sum(), 2)

                tacos = round((day_spends / ov_sales) * 100, 2) if ov_sales > 0 else 0.0
                troas = round(ov_sales / day_spends, 2) if day_spends > 0 else 0.0

                rows.append({
                    'Date (Day)': f"Day {day_i}",
                    'Total Sales': ov_sales,
                    'Total Orders': ov_orders,
                    'Total Spends': day_spends,
                    'TACOS': f"{tacos:.2f}%",
                    'TROAS': troas
                })

            res_df = pd.DataFrame(rows)
            total_sales_sum = res_df['Total Sales'].sum()
            total_orders_sum = res_df['Total Orders'].sum()
            total_spends_sum = res_df['Total Spends'].sum()
            avg_tacos = round((total_spends_sum / total_sales_sum) * 100, 2) if total_sales_sum > 0 else 0.0
            avg_troas = round(total_sales_sum / total_spends_sum, 2) if total_spends_sum > 0 else 0.0

            summary_row = pd.DataFrame([{
                'Date (Day)': 'Aggregate / Average',
                'Total Sales': total_sales_sum,
                'Total Orders': total_orders_sum,
                'Total Spends': total_spends_sum,
                'TACOS': f"{avg_tacos:.2f}%",
                'TROAS': avg_troas
            }])
            res_df = pd.concat([res_df, summary_row], ignore_index=True)
            return res_df

        def _format_dashboard_value(value, column_name=""):
            if pd.isna(value):
                return value
            col = str(column_name).strip().upper()
            is_percentage = ('%' in col) or ('ACOS' in col) or ('TACOS' in col)

            if isinstance(value, (int, np.integer)) and not isinstance(value, bool):
                return f"{int(value):,}" if not is_percentage else f"{int(value)}%"

            if isinstance(value, (float, np.floating)) and np.isfinite(value):
                if is_percentage:
                    return f"{value:.2f}%"
                if float(value).is_integer():
                    return f"{int(value):,}"
                return f"{value:,.2f}"
            return value

        def format_dashboard_dataframe(df):
            out = df.copy()
            for col in out.columns:
                out[col] = out[col].map(lambda v: _format_dashboard_value(v, col))
            return out

        def render_unified_single_table(df_to_show, key_prefix="mom"):
            if df_to_show is None or df_to_show.empty:
                return
            display_df = format_dashboard_dataframe(df_to_show)
            st.dataframe(
                display_df,
                use_container_width=True,
                hide_index=False if (isinstance(df_to_show.index, pd.Index) and df_to_show.index.name) else True,
                key=f"{key_prefix}_grid"
            )

        st.markdown("### 🔍 Global Dashboard Filters")
        available_months_asc = get_calendar_months_asc(final_df)
        selected_month = st.selectbox("Select Month Across Dashboard (Excluding Comparison Tables)", ["All Months"] + available_months_asc)

        filtered_df = final_df.copy()
        if selected_month != "All Months":
            filtered_df = filtered_df[filtered_df['Month'] == selected_month]

        total_ad_sales = round(filtered_df['_sales'].sum(), 2)
        total_ad_spends = round(filtered_df['_budget_consumed'].sum(), 2)
        overall_ad_roas = round(total_ad_sales / total_ad_spends, 2) if total_ad_spends > 0 else 0.0

        if not master_sales_df.empty and '_sales_dt' in master_sales_df.columns:
            f_sales_df = master_sales_df if selected_month == "All Months" else master_sales_df[master_sales_df['Month'] == selected_month]
            total_overall_sales = round(f_sales_df['_gross_amt'].sum(), 2)
            total_overall_orders = len(f_sales_df)
        else:
            total_overall_sales = total_ad_sales
            total_overall_orders = int(filtered_df['_orders'].sum())

        overall_tacos = round((total_ad_spends / total_overall_sales) * 100, 2) if total_overall_sales > 0 else 0.0
        overall_troas = round(total_overall_sales / total_ad_spends, 2) if total_ad_spends > 0 else 0.0

        st.markdown("### 📈 Overall Campaign Performance Dashboard")
        
        dcol1, dcol2, dcol3 = st.columns(3)
        with dcol1: st.metric("Ad Sales", f"₹{total_ad_sales:,.2f}")
        with dcol2: st.metric("Total Sales", f"₹{total_overall_sales:,.2f}")
        with dcol3: st.metric("Ad ROAS", f"{overall_ad_roas:.2f}x")

        dcol4, dcol5, dcol6 = st.columns(3)
        with dcol4: st.metric("TACOS", f"{overall_tacos:.2f}%")
        with dcol5: st.metric("TROAS", f"{overall_troas:.2f}")
        with dcol6: st.metric("Total Orders", f"{int(total_overall_orders):,}")

        st.divider()

        st.markdown("### 📑 Navigation & Performance Breakdown")
        
        main_tab1, main_tab2, main_tab3, main_tab4, main_tab5, main_tab6, main_tab7, main_tab8, main_tab9 = st.tabs([
            "📄 Raw Files Preview",
            "📌 Consolidated Master",
            "📊 Comparison Tables (MoM)",
            "📅 Daily Tracker",
            "📈 Interactive Trend Analytics",
            "🎯 Campaign Performance",
            "📢 Ad Type Performance",
            "🔎 Keyword Performance",
            "📅 Weekly Trend"
        ])

        with main_tab1:
            st.caption("Inspect individual sheets tab-by-tab for each uploaded file.")
            selected_file_name = st.selectbox("Select Uploaded File to Preview:", list(raw_files_dict.keys()))
            if selected_file_name:
                sheets = raw_files_dict[selected_file_name]
                selected_sheet = st.selectbox("Select Sheet Tab:", list(sheets.keys()))
                if selected_sheet:
                    st.dataframe(format_dashboard_dataframe(sheets[selected_sheet].head(100)), use_container_width=True, hide_index=True)

        with main_tab2:
            st.caption("Preview the combined dataset across all uploaded files.")
            preview_clean_df = final_df.drop(columns=['_impressions', '_direct_atc', '_indirect_atc', '_atc', '_direct_orders', '_indirect_orders', '_orders', '_direct_sales', '_indirect_sales', '_sales', '_budget_consumed', '_date_dt', 'Ad Type Combined'], errors='ignore')
            st.dataframe(format_dashboard_dataframe(preview_clean_df.head(100)), use_container_width=True, hide_index=True)

        with main_tab3:
            st.subheader("📊 Month-on-Month Comparison Tables")
            comp_sub_tab0, comp_sub_tab1, comp_sub_tab2, comp_sub_tab3, comp_sub_tab4 = st.tabs([
                "📅 Monthly Summary",
                "🎯 Campaign Comparison",
                "📢 Ad Type Comparison",
                "🔎 Keyword Comparison",
                "📅 Weekly Comparison"
            ])

            with comp_sub_tab0:
                st.markdown("#### Monthly Summary Table (with Sales, Orders, TACOS, TROAS)")
                monthly_summary_df = create_monthly_summary_table(final_df)
                if not monthly_summary_df.empty:
                    render_unified_single_table(monthly_summary_df, key_prefix="monthly_summary_tab")
                    excel_month_summary = style_and_export_pivot(monthly_summary_df, sheet_name="Monthly_Summary")
                    st.download_button(
                        label="📥 Download Monthly Summary Table (.xlsx)",
                        data=excel_month_summary,
                        file_name="Monthly_Summary_Report.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="btn_dl_month_summary"
                    )

            with comp_sub_tab1:
                camp_pivot = create_mom_comparison_table(final_df, 'Campaign Name') if 'Campaign Name' in final_df.columns else None
                if camp_pivot is not None and not camp_pivot.empty:
                    render_unified_single_table(camp_pivot, key_prefix="camp_pivot_tab")

            with comp_sub_tab2:
                ad_pivot = create_mom_comparison_table(final_df, 'Ad Type Combined') if 'Ad Type Combined' in final_df.columns else None
                if ad_pivot is not None and not ad_pivot.empty:
                    render_unified_single_table(ad_pivot, key_prefix="ad_pivot_tab")

            with comp_sub_tab3:
                kw_col = next((c for c in ['Search Term', 'Keyword', 'Targeting Value'] if c in final_df.columns), None)
                kw_pivot = create_mom_comparison_table(final_df, kw_col) if kw_col else None
                if kw_pivot is not None and not kw_pivot.empty:
                    render_unified_single_table(kw_pivot, key_prefix="kw_pivot_tab")

            with comp_sub_tab4:
                weekly_comp_df = create_weekly_comparison_table_with_sales(final_df)
                if not weekly_comp_df.empty:
                    render_unified_single_table(weekly_comp_df, key_prefix="weekly_comp_tab")

        with main_tab4:
            st.subheader("📅 Daily Tracker")
            tracker_month = st.selectbox("Select Month for Daily Tracker:", available_months_asc, key="tracker_month_sel")
            if tracker_month:
                daily_df = create_daily_tracker_table(tracker_month)
                if not daily_df.empty:
                    render_unified_single_table(daily_df, key_prefix="daily_tracker_tab")
                    excel_daily = style_and_export_pivot(daily_df, sheet_name="Daily_Tracker")
                    st.download_button(
                        label="📥 Download Daily Tracker Table (.xlsx)",
                        data=excel_daily,
                        file_name=f"Daily_Tracker_{tracker_month}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="btn_dl_daily_tracker"
                    )
                else:
                    st.info("No sales data available for the selected month.")

        with main_tab5:
            st.subheader("📈 Interactive Multi-Metric Trend Analytics")
            all_df_months = get_calendar_months_asc(final_df)
            selected_trend_months = st.multiselect("Select Months:", options=all_df_months, default=all_df_months, key="trend_m_sel")
            if selected_trend_months:
                trend_df = final_df[final_df['Month'].isin(selected_trend_months)].copy()
                monthly_summary = trend_df.groupby('Month').agg(
                    _impressions=('_impressions', 'sum'),
                    _sales=('_sales', 'sum'),
                    _budget_consumed=('_budget_consumed', 'sum')
                ).reset_index()
                st.dataframe(format_dashboard_dataframe(monthly_summary), use_container_width=True, hide_index=True)

        with main_tab6:
            if 'Campaign Name' in filtered_df.columns:
                st.markdown("#### 🎯 Campaign Performance")
                camp_df = compute_grouped_table(filtered_df, 'Campaign Name', "All")
                st.dataframe(format_dashboard_dataframe(camp_df), use_container_width=True, hide_index=True)

        with main_tab7:
            if 'Ad Type Combined' in filtered_df.columns:
                st.markdown("#### 📢 Ad Type Performance")
                ad_df = compute_grouped_table(filtered_df, 'Ad Type Combined', "All")
                st.dataframe(format_dashboard_dataframe(ad_df), use_container_width=True, hide_index=True)

        with main_tab8:
            kw_col = next((c for c in ['Search Term', 'Keyword', 'Targeting Value'] if c in final_df.columns), None)
            if kw_col:
                st.markdown("#### 🔎 Keyword Performance")
                kw_df = compute_grouped_table(filtered_df, kw_col, "All")
                st.dataframe(format_dashboard_dataframe(kw_df), use_container_width=True, hide_index=True)

        with main_tab9:
            if 'Week' in filtered_df.columns and filtered_df['Week'].notna().any():
                st.markdown("#### 📅 Weekly Trend")
                wk_df = compute_grouped_table(filtered_df, 'Week', "All")
                st.dataframe(format_dashboard_dataframe(wk_df), use_container_width=True, hide_index=True)

    else:
        st.info("Please upload valid ad campaign reports to start analysis.")
else:
    st.info("Awaiting file uploads. Please upload your ad campaign spreadsheets and sales reports.")
